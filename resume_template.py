"""
User-uploaded resume templates.

A user uploads a .docx whose layout they like (e.g. a SheetsResume-style sample).
We learn one "prototype" paragraph per resume role (name, contact, section
heading, company line, title line, bullet, sub-bullet, spacer) from that file,
then build tailored resumes by cloning those paragraphs' formatting — fonts,
borders, numbering, spacing — and filling in the tailored text. The template's
own sample content is discarded.

Storage: data/users/<user_id>/resume_template.docx (+ resume_template.json).
"""

import copy
import json
import logging
import re
import zipfile
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Twips
from docx.text.paragraph import Paragraph

import storage

logger = logging.getLogger(__name__)

MAX_TEMPLATE_BYTES = 5 * 1024 * 1024

DATE_PATTERN = r"(Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec|Present|\d{4})"
KNOWN_SECTIONS = {
    "SUMMARY", "CORE CAPABILITIES", "PROFESSIONAL EXPERIENCE",
    "EARLIER EXPERIENCE", "TECHNICAL SKILLS", "CERTIFICATIONS",
    "EDUCATION", "CERTIFICATIONS & EDUCATION", "SELECTED INNOVATION",
    "SELECTED SYSTEMS & INNOVATION", "PROJECT HIGHLIGHTS", "KEY SKILLS",
}

SAMPLE_RESUME = """JANE SAMPLE
jane@example.com · (555) 555-0100 · Washington, DC · linkedin.com/in/janesample

SUMMARY
Data and analytics leader with 12 years of experience building BI teams.

PROFESSIONAL EXPERIENCE
Director of Analytics
Example Corp · Jan 2020 – Present
• Led a 15-person analytics organization supporting $2B in annual revenue.
  - Cut reporting turnaround from 5 days to 1 day.
• Built the enterprise data platform on Snowflake and dbt.

Senior Manager, Business Intelligence
Sample Industries · Mar 2015 – Dec 2019
• Delivered self-service dashboards adopted by 400+ users.

TECHNICAL SKILLS
Languages: SQL; Python; R
Platforms: Snowflake; Tableau; Power BI

EDUCATION
M.S., Information Systems · Example University · 2012
"""


# ── Storage ────────────────────────────────────────────────────────────────────

def template_path(user_id: str) -> Path:
    return storage.get_user_dir(user_id) / "resume_template.docx"


def _meta_path(user_id: str) -> Path:
    return storage.get_user_dir(user_id) / "resume_template.json"


def has_template(user_id: str) -> bool:
    return template_path(user_id).exists()


def template_info(user_id: str) -> dict:
    if not has_template(user_id):
        return {"loaded": False}
    meta = {}
    try:
        meta = json.loads(_meta_path(user_id).read_text())
    except Exception:
        pass
    return {"loaded": True, **meta}


def delete_template(user_id: str) -> None:
    template_path(user_id).unlink(missing_ok=True)
    _meta_path(user_id).unlink(missing_ok=True)


def save_template(user_id: str, data: bytes, filename: str) -> dict:
    """Validate and store an uploaded template. Raises ValueError with a user-facing message."""
    if not filename.lower().endswith(".docx"):
        raise ValueError("Only .docx files are supported")
    if len(data) > MAX_TEMPLATE_BYTES:
        raise ValueError("Template is too large (5 MB max)")
    if not data.startswith(b"PK"):
        raise ValueError("That file is not a valid .docx")

    storage.ensure_user_dirs(user_id)
    tmp = template_path(user_id).with_suffix(".upload")
    tmp.write_bytes(data)
    try:
        with zipfile.ZipFile(tmp) as z:
            names = z.namelist()
            if "word/document.xml" not in names:
                raise ValueError("That file is not a valid .docx")
            if any(n.lower().endswith("vbaproject.bin") for n in names):
                raise ValueError("Macro-enabled documents are not allowed")
        protos = analyze(tmp)
    except zipfile.BadZipFile:
        tmp.unlink(missing_ok=True)
        raise ValueError("That file is not a valid .docx")
    except Exception:
        tmp.unlink(missing_ok=True)
        raise
    tmp.replace(template_path(user_id))

    info = {
        "filename": Path(filename).name,
        "uploaded_at": datetime.now(timezone.utc).isoformat(),
        "roles": sorted(protos.keys()),
    }
    _meta_path(user_id).write_text(json.dumps(info))
    return {"loaded": True, **info}


# ── Template analysis ──────────────────────────────────────────────────────────

def _numbering(p):
    """(numId, ilvl) for a numbered/bulleted paragraph, else None."""
    ppr = p._p.pPr
    if ppr is None or ppr.numPr is None:
        return None
    ilvl = ppr.numPr.ilvl.val if ppr.numPr.ilvl is not None else 0
    return ppr.numPr.numId.val, ilvl


def _has_bottom_border(p) -> bool:
    ppr = p._p.pPr
    return ppr is not None and ppr.find(qn("w:pBdr")) is not None


def _has_tab(p) -> bool:
    return any(r._r.findall(qn("w:tab")) for r in p.runs)


def _first_bold(p) -> bool:
    runs = [r for r in p.runs if r.text.strip()]
    return bool(runs) and bool(runs[0].bold)


def analyze(path) -> dict:
    """Find one prototype paragraph per resume role. Raises ValueError if unusable."""
    doc = Document(str(path))
    paras = doc.paragraphs
    filled = [p for p in paras if p.text.strip()]
    if len(filled) < 4:
        raise ValueError("That template has too little content to learn a layout from")

    protos = {"name": filled[0], "contact": filled[1]}
    for p in paras:
        if not p.text.strip():
            if "spacer" not in protos and p._p.pPr is not None:
                protos["spacer"] = p
            continue
        if p is filled[0] or p is filled[1]:
            continue
        num = _numbering(p)
        if num:
            key = "sub" if num[1] >= 1 else "bullet"
        elif _has_bottom_border(p):
            key = "section"
        elif _has_tab(p) and _first_bold(p):
            key = "org"
        elif _has_tab(p):
            key = "title"
        else:
            key = None
        if key and key not in protos:
            protos[key] = p

    if "section" not in protos:
        raise ValueError(
            "Couldn't find a section heading in that template "
            "(expected a heading with an underline/border, e.g. WORK EXPERIENCE)"
        )
    if "bullet" not in protos:
        raise ValueError("Couldn't find a bulleted line in that template")
    return protos


def _text_width_twips(doc) -> int:
    sec = doc.sections[0]
    return int((sec.page_width - sec.left_margin - sec.right_margin) / 635)  # EMU → twips


def _rpr_of(p, plain: bool = False):
    """rPr of a prototype's first direct text run (preferring a non-bold run if plain)."""
    runs = [r for r in p.runs if r.text.strip()] or list(p.runs)
    if plain:
        non_bold = [r for r in runs if not r.bold]
        runs = non_bold or runs
    for r in runs:
        if r._r.rPr is not None:
            return copy.deepcopy(r._r.rPr)
    ppr = p._p.pPr
    if ppr is not None and ppr.find(qn("w:rPr")) is not None:
        return copy.deepcopy(ppr.find(qn("w:rPr")))
    return None


# ── Text → blocks ──────────────────────────────────────────────────────────────
# Mirrors the line heuristics in docx_builder.generate_resume_docx for the older
# "Title / Org · dates" shape, and also understands the three-line job header
# the tailoring step emits ("Title", "Org", "2021 – Present" in either order).

_MONTH = r"(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Sept|Oct|Nov|Dec)[a-z]*\.?"
DATE_RANGE = re.compile(
    rf"^(?:{_MONTH}\s+)?\d{{4}}\s*[–—-]\s*(?:(?:{_MONTH}\s+)?\d{{4}}|Present|Current)$", re.I
)
TITLE_WORDS = re.compile(
    r"\b(manager|analyst|architect|director|consultant|consulting|founder|co-founder|engineer|lead|"
    r"partner|strategist|officer|developer|scientist|specialist|president|ceo|cto|cfo|vp|head|"
    r"associate|administrator|coordinator|designer|advisor|owner|principal|intern|executive)\b",
    re.I,
)


def _split_org_title(a: str, b: str) -> tuple[str, str]:
    """Two header lines in unknown order -> (org, title)."""
    ta, tb = bool(TITLE_WORDS.search(a)), bool(TITLE_WORDS.search(b))
    if ta and not tb:
        return b, a
    if tb and not ta:
        return a, b
    if ta and tb:
        if a.isupper() and not b.isupper():
            return a, b
        if b.isupper() and not a.isupper():
            return b, a
    return b, a


def _job_header(lines: list[str], i: int):
    """If lines[i:] opens with 'L1 [, L2], DATE-RANGE' return (title, org, dates, lines_used)."""
    group = []
    for j in range(i, min(i + 3, len(lines))):
        s = lines[j].strip()
        if not s or s[0] in "•-*" or len(s) > 90:
            return None
        group.append(s)
        if DATE_RANGE.match(s):
            break
    else:
        return None
    if len(group) < 2 or not DATE_RANGE.match(group[-1]):
        return None
    first = group[0]
    if first.upper().rstrip(":") in KNOWN_SECTIONS or DATE_RANGE.match(first):
        return None
    dates = group[-1]
    if len(group) == 2:
        return "", first, dates, 2
    org, title = _split_org_title(first, group[1])
    return title, org, dates, 3


# "Title | Company    Sep 2021 - Present": a header line that ends in a date range
DATE_TAIL = re.compile(
    rf"(?:^|\s)((?:{_MONTH}\s+)?\d{{4}}\s*[–—-]\s*(?:(?:{_MONTH}\s+)?\d{{4}}|Present|Current))$", re.I
)
PHONE = re.compile(r"\(?\d{3}\)?[-.\s]\d{3}[-.\s]\d{4}")
DOMAIN = re.compile(r"\b[\w-]+\.(?:com|dev|io|net|org|me|co|ai|us|edu)\b|linkedin|github|https?://", re.I)


def _is_contact_line(line: str) -> bool:
    return "@" in line or bool(PHONE.search(line)) or bool(DOMAIN.search(line))


def _dated_header(line: str):
    """'Title | Company    Sep 2021 - Present' -> (title, org, dates), else None."""
    m = DATE_TAIL.search(line)
    if not m or len(line) > 150:
        return None
    left = line[:m.start()].rstrip(" \t|·–—,")
    if not left or (" | " not in left and len(left.split()) > 12):
        return None  # a sentence that happens to end in a date range
    if " | " in left:
        title, org = (x.strip() for x in left.split(" | ", 1))
        return title, org, m.group(1)
    return "", left, m.group(1)


def _is_location(line: str) -> bool:
    if len(line) > 45 or line.endswith(".") or len(line.split()) > 6 or DATE_TAIL.search(line):
        return False
    return "," in line or bool(re.search(r"\b(area|remote|hybrid|on-?site)$", line, re.I))


def parse_blocks(text: str) -> list[tuple]:
    lines = text.strip().split("\n")
    blocks: list[tuple] = []
    i = 0
    in_header = True
    in_job = False
    bullet_base = None  # indent of the first bullet in the current run

    def next_nonempty(k):
        while k < len(lines) and not lines[k].strip():
            k += 1
        return lines[k] if k < len(lines) else ""

    while i < len(lines):
        raw = lines[i]
        line = raw.strip()
        if not line:
            if blocks:
                in_header = False
            i += 1
            continue

        if not blocks:
            blocks.append(("name", line))
            i += 1
            continue

        # Header region: everything between the name and the first blank line / section heading
        if in_header:
            if line.isupper() and len(line) > 3 and not _is_contact_line(line):
                in_header = False
            elif _is_contact_line(line):
                blocks.append(("contact", line))
                i += 1
                continue
            else:
                blocks.append(("tagline", line))
                i += 1
                continue

        if line[0] in "•-*":
            indent = len(raw) - len(raw.lstrip())
            if bullet_base is None:
                bullet_base = indent
            kind = "sub" if indent >= bullet_base + 2 else "bullet"
            blocks.append((kind, line.lstrip("•-* ").strip()))
            i += 1
            continue
        bullet_base = None

        dated = _dated_header(line)
        if dated:
            title, org, dates = dated
            location = ""
            nxt = lines[i + 1].strip() if i + 1 < len(lines) else ""
            if nxt and _is_location(nxt):
                location = nxt
                i += 1
            blocks.append(("job", title, org, dates, location))
            in_job = True
            i += 1
            continue

        header = _job_header(lines, i)
        if header:
            title, org, dates, used = header
            blocks.append(("job", title, org, dates))
            in_job = True
            i += used
            continue

        is_section = (
            (line.isupper() and len(line) > 3)
            or line.upper().rstrip(":") in KNOWN_SECTIONS
        )
        if is_section:
            blocks.append(("section", line.rstrip(":")))
            in_job = False
            i += 1
            continue

        if re.match(r"^(BI |Senior |Operations |IT |Co-Founder|Product )", line) and "(" in line:
            m = re.match(r"^(.*?)\s*\((.+)\)\s*$", line)
            blocks.append(("subheader", m.group(1), m.group(2)) if m else ("subheader", line, ""))
            i += 1
            continue

        next_line = lines[i + 1].strip() if i + 1 < len(lines) else ""
        looks_like_title = (
            len(line) < 80
            and next_line
            and ("–" in next_line or "—" in next_line or "·" in next_line or re.match(r"^[A-Z]", next_line))
        )
        if re.search(DATE_PATTERN, next_line) and looks_like_title:
            org, dates = next_line, ""
            if "\t" in next_line:
                parts = next_line.split("\t")
                org, dates = parts[0].strip(), parts[-1].strip() if len(parts) > 1 else ""
            else:
                m = re.search(r"((?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec|Present)[\w\s,–\-]*\d{4}.*$)", next_line)
                if m:
                    org, dates = next_line[:m.start()].rstrip(" ·–—,"), m.group(1).strip()
            blocks.append(("job", line, org, dates))
            in_job = True
            i += 2
            continue

        # "Category: description" (skills, competencies) -> bullet with a bold label
        head = line.split(":", 1)[0]
        if ":" in line and len(head) < 45 and len(head.split()) <= 6 and not re.search(r"[.!?]", head):
            label, value = line.split(":", 1)
            blocks.append(("skill", label, value.strip()))
            i += 1
            continue

        # Short heading directly above bullets (e.g. a skills category)
        nxt = next_nonempty(i + 1).strip()
        if len(line) < 60 and not line.endswith(".") and nxt[:1] in ("•", "-", "*") and not in_job:
            blocks.append(("label", line))
        elif in_job:
            blocks.append(("bullet", line))
        else:
            blocks.append(("body", line))
        i += 1
    return blocks


# ── Contact line from profile ─────────────────────────────────────────────────

def _normalize_url(url: str) -> str:
    url = url.strip()
    return url if re.match(r"^[a-z][a-z0-9+.-]*:", url, re.I) else "https://" + url


def _display_url(url: str) -> str:
    return re.sub(r"^(?:https?://)?(?:www\.)?", "", url.strip(), flags=re.I).rstrip("/")


def contact_segments(config) -> list[tuple[str, str | None]]:
    """Contact line pieces from the profile as (text, url-or-None); empty if nothing is set."""
    segs: list[tuple[str, str | None]] = []
    loc = ", ".join(x for x in (config.author_city, config.author_state) if x) or config.author_location
    if loc:
        segs.append((loc, None))
    if config.author_email:
        segs.append((config.author_email, "mailto:" + config.author_email))
    if config.author_phone:
        segs.append((config.author_phone, None))
    for url in (config.author_linkedin, config.author_github, config.author_website, config.author_portfolio):
        if url and url.strip():
            segs.append((_display_url(url), _normalize_url(url)))
    return segs


def _add_hyperlink(p: Paragraph, url: str, text: str, fmt):
    """Append an external hyperlink run to p; fmt(run) applies the base formatting."""
    from docx.oxml import OxmlElement
    from docx.shared import RGBColor
    from docx.text.run import Run

    r_id = p.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    r = OxmlElement("w:r")
    link.append(r)
    p._p.append(link)
    run = Run(r, p)
    fmt(run)
    run.text = text
    run.font.color.rgb = RGBColor(0x05, 0x63, 0xC1)
    run.font.underline = True
    return run


def add_contact_runs(p: Paragraph, segments, fmt, sep: str = " | "):
    """Write contact segments into p, hyperlinking the ones that have a URL."""
    for n, (text, url) in enumerate(segments):
        if n:
            fmt(p.add_run(sep))
        if url:
            _add_hyperlink(p, url, text, fmt)
        else:
            run = p.add_run(text)
            fmt(run)


# ── Rendering ──────────────────────────────────────────────────────────────────

class _Builder:
    def __init__(self, template: Path):
        self.doc = Document(str(template))
        self.protos = analyze(template)
        body = self.doc.element.body
        for child in list(body):
            if child.tag != qn("w:sectPr"):
                body.remove(child)
        self._sect = body.find(qn("w:sectPr"))
        self.width = _text_width_twips(self.doc)
        self._rpr = {k: _rpr_of(p, plain=k in ("bullet", "sub")) for k, p in self.protos.items()}

    def _fallback(self, role: str) -> str:
        chain = {"sub": "bullet", "org": "title", "title": "org", "contact": "name",
                 "spacer": None}
        while role not in self.protos:
            role = chain.get(role) or "bullet"
        return role

    def _new(self, role: str, right_tab: bool = False, bulleted: bool = True) -> Paragraph:
        role = self._fallback(role)
        proto = self.protos[role]
        el = copy.deepcopy(proto._p)
        for attr in [a for a in el.attrib if a.endswith("}paraId") or a.endswith("}textId")]:
            del el.attrib[attr]
        for child in list(el):
            if child.tag != qn("w:pPr"):
                el.remove(child)
        self._sect.addprevious(el)
        p = Paragraph(el, self.doc._body)
        if not bulleted and p._p.pPr is not None and p._p.pPr.numPr is not None:
            p._p.pPr.remove(p._p.pPr.numPr)
        if right_tab:
            p.paragraph_format.tab_stops.clear_all()
            p.paragraph_format.tab_stops.add_tab_stop(Twips(self.width), WD_TAB_ALIGNMENT.RIGHT)
        p._role = role
        return p

    def _fmt(self, p: Paragraph, bold=None, italic=None):
        """Returns fmt(run): the prototype's run formatting for p's role, plus overrides."""
        rpr = self._rpr.get(p._role)

        def fmt(run):
            if rpr is not None:
                run._r.insert(0, copy.deepcopy(rpr))
            if bold is not None:
                run.bold = bold
            if italic is not None:
                run.italic = italic

        return fmt

    def _run(self, p: Paragraph, text: str, bold=None, italic=None):
        run = p.add_run(text)
        self._fmt(p, bold, italic)(run)
        return run

    def spacer(self):
        if "spacer" in self.protos:
            self._new("spacer")

    def add(self, block: tuple):
        kind = block[0]
        if kind in ("name", "section"):
            self._run(self._new(kind), block[1])
        elif kind == "contact":
            text = re.sub(r"\s*[·•]\s*", " | ", block[1])
            self._run(self._new("contact"), text)
        elif kind == "contact_links":
            p = self._new("contact")
            fmt = self._fmt(p)
            details = [sg for sg in block[1] if not sg[1] or sg[1].startswith("mailto:")]
            links = [sg for sg in block[1] if sg[1] and not sg[1].startswith("mailto:")]
            add_contact_runs(p, details, fmt)
            if details and links:
                br = p.add_run()
                fmt(br)
                br.add_break()
            add_contact_runs(p, links, fmt)
        elif kind == "body":
            self._run(self._new("bullet", bulleted=False), block[1])
        elif kind == "label":
            self._run(self._new("bullet", bulleted=False), block[1], bold=True)
        elif kind == "tagline":
            self._run(self._new("title"), block[1], italic=True)
        elif kind in ("bullet", "sub"):
            self._run(self._new(kind), block[1])
        elif kind == "skill":
            p = self._new("bullet")
            self._run(p, block[1] + ": ", bold=True)
            self._run(p, block[2], bold=False)
        elif kind == "job":
            title, org, dates = block[1:4]
            location = block[4] if len(block) > 4 else ""
            p = self._new("org", right_tab=True)
            self._run(p, org, bold=True)
            if dates:
                self._run(p, "\t", bold=True)
                self._run(p, dates, bold=True)
            if title or location:
                p = self._new("title", right_tab=bool(location))
                self._run(p, title, italic=True)
                if location:
                    self._run(p, "\t", italic=True)
                    self._run(p, location, italic=True)
        elif kind == "subheader":
            _, title, dates = block
            p = self._new("title", right_tab=True)
            self._run(p, title, italic=True)
            if dates:
                self._run(p, "\t", italic=True)
                self._run(p, dates, italic=True)


def apply_profile_contact(blocks: list[tuple], segments) -> list[tuple]:
    """Replace the model's contact line with one built from the profile (so links are real)."""
    if not segments:
        return blocks
    blocks = [b for b in blocks if b[0] != "contact"]
    at = 1 if blocks and blocks[0][0] == "name" else 0
    return blocks[:at] + [("contact_links", segments)] + blocks[at:]


def render(tailored_text: str, user_id: str, out_path: Path) -> str:
    """Build a resume from the user's template. Returns out_path."""
    b = _Builder(template_path(user_id))
    blocks = parse_blocks(tailored_text)
    config = storage.load_config(user_id)
    blocks = apply_profile_contact(blocks, contact_segments(config))
    if config.author_name and blocks and blocks[0][0] == "name":
        blocks[0] = ("name", config.author_name)
    prev = None
    for block in blocks:
        kind = block[0]
        if kind == "section" and prev is not None:
            b.spacer()
        elif kind == "job" and prev in ("bullet", "sub", "skill", "job", "body"):
            b.spacer()
        b.add(block)
        prev = kind
    b.doc.save(str(out_path))
    _check_relationships(out_path)
    return str(out_path)


def _check_relationships(path: Path) -> None:
    """Raise if document.xml references a relationship id that isn't in its .rels (Word calls that corrupt)."""
    with zipfile.ZipFile(path) as z:
        used = set(re.findall(r'r:(?:id|embed|link)="([^"]+)"', z.read("word/document.xml").decode("utf8")))
        defined = set(re.findall(r'Id="([^"]+)"', z.read("word/_rels/document.xml.rels").decode("utf8")))
    missing = used - defined
    if missing:
        raise ValueError(f"unresolved relationships in generated resume: {sorted(missing)}")


def render_sample(user_id: str, out_path: Path) -> str:
    return render(SAMPLE_RESUME, user_id, out_path)
